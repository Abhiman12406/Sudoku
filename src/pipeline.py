"""End-to-end walking skeleton and hierarchical pipeline orchestrator."""

import io
from pathlib import Path
from typing import (
    Any,
    BinaryIO,
    Dict,
    Iterator,
    List,
    Literal,
    Optional,
    Sequence,
    Tuple,
    Union,
    overload,
)

from src.domain.models import AllocatedSignal, IOType, SegregationPolicy, Signal, TechnicalQuery
from src.export.excel_export import export_allocation_to_excel
from src.ingestion.parser import parse_and_validate_signals
from src.solver.audit import verify_allocation_invariants
from src.solver.hierarchical_solver import solve_hierarchical_packing
from src.solver.walking_solver import solve_walking_skeleton


class PipelineResult:
    """Container holding allocations, deliverable Excel bytes, and generated TQs."""

    def __init__(
        self,
        allocations: List[AllocatedSignal],
        excel_bytes: bytes,
        tqs: List[TechnicalQuery],
    ):
        self.allocations = allocations
        self.excel_bytes = excel_bytes
        self.tqs = tqs

    def __iter__(self) -> Any:
        # Enables backwards-compatible 2-element tuple unpacking:
        # allocations, out_bytes = run_walking_pipeline(...)
        return iter((self.allocations, self.excel_bytes))

    def __getitem__(self, index):
        return (self.allocations, self.excel_bytes, self.tqs)[index]

    def __len__(self):
        return 2

    def __repr__(self) -> str:
        return f"<PipelineResult: {len(self.allocations)} allocations, {len(self.tqs)} TQs, {len(self.excel_bytes)} bytes>"


class RevisionPipelineResult:
    """Container holding revision allocations, deliverable Excel bytes with Revision_Impact tab, and diff metrics."""

    def __init__(
        self,
        allocations: List[AllocatedSignal],
        excel_bytes: bytes,
        revision_result: Any,
        tqs_a: Optional[List[TechnicalQuery]] = None,
        tqs_b: Optional[List[TechnicalQuery]] = None,
    ):
        self.allocations = allocations
        self.excel_bytes = excel_bytes
        self.revision_result = revision_result
        self.tqs_a = tqs_a or []
        self.tqs_b = tqs_b or []
        self.tqs = (tqs_a or []) + (tqs_b or [])

    def __iter__(self):
        return iter((self.allocations, self.excel_bytes, self.revision_result))

    def __getitem__(self, index):
        return (self.allocations, self.excel_bytes, self.revision_result)[index]

    def __len__(self):
        return 3

    def __repr__(self) -> str:
        score = getattr(self.revision_result, "perturbation_score", 0)
        return f"<RevisionPipelineResult: {len(self.allocations)} allocations, score={score}>"


@overload
def run_walking_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: Literal[False] = False,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> PipelineResult:
    ...


@overload
def run_walking_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: Literal[True] = ...,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> Tuple[List[AllocatedSignal], bytes, List[TechnicalQuery]]:
    ...


@overload
def run_walking_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: bool = ...,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> Union[PipelineResult, Tuple[List[AllocatedSignal], bytes, List[TechnicalQuery]]]:
    ...


def run_walking_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: bool = False,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> Union[PipelineResult, Tuple[List[AllocatedSignal], bytes, List[TechnicalQuery]]]:
    """Execute end-to-end ingestion, fault-tolerant TQ isolation, CP-SAT allocation, and Excel export.

    Args:
        file_or_path_or_bytes: Path to input spreadsheet, binary buffer, or raw bytes.
        charms_per_baseplate: Number of CHARMs per baseplate (default: 12).
        catalog: Optional custom hardware catalog mapping.
        return_tqs: If True, returns explicit 3-tuple (allocations, excel_bytes, tqs).
        template_file_or_bytes: Optional blank customer Excel template to overlay onto.

    Returns:
        PipelineResult (unpackable as (allocations, excel_bytes) or providing .tqs).
    """
    if isinstance(file_or_path_or_bytes, bytes):
        buffer: Union[str, Path, BinaryIO] = io.BytesIO(file_or_path_or_bytes)
    else:
        buffer = file_or_path_or_bytes

    # 1. Parse and validate with fault isolation
    signals, tqs = parse_and_validate_signals(buffer)

    # 2. Solve CP-SAT slot assignment for valid signals
    if signals:
        allocations = solve_walking_skeleton(
            signals,
            charms_per_baseplate=charms_per_baseplate,
            catalog=catalog,
        )
        valid, errors = verify_allocation_invariants(
            allocations,
            charms_per_baseplate=charms_per_baseplate,
        )
        if not valid:
            raise RuntimeError(f"Allocation violated physical constraints: {'; '.join(errors)}")
    else:
        allocations = []

    # 3. Export populated Excel workbook with dedicated TQ_Log
    excel_bytes = export_allocation_to_excel(
        allocations,
        tqs=tqs,
        template_path_or_buffer=template_file_or_bytes,
        charms_per_bp=charms_per_baseplate,
    )

    if return_tqs:
        return allocations, excel_bytes, tqs

    return PipelineResult(allocations, excel_bytes, tqs)


@overload
def run_hierarchical_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    baseplates_per_cioc: int = 8,
    spare_margin: float = 0.0,
    declared_sizes: Optional[dict[str, int]] = None,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: Literal[False] = False,
    segregation_policy: Union[str, SegregationPolicy] = SegregationPolicy.CIOC,
    ciocs_per_cabinet: int = 4,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> PipelineResult:
    ...


@overload
def run_hierarchical_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    baseplates_per_cioc: int = 8,
    spare_margin: float = 0.0,
    declared_sizes: Optional[dict[str, int]] = None,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: Literal[True] = ...,
    segregation_policy: Union[str, SegregationPolicy] = SegregationPolicy.CIOC,
    ciocs_per_cabinet: int = 4,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> Tuple[List[AllocatedSignal], bytes, List[TechnicalQuery]]:
    ...


@overload
def run_hierarchical_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    baseplates_per_cioc: int = 8,
    spare_margin: float = 0.0,
    declared_sizes: Optional[dict[str, int]] = None,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: bool = ...,
    segregation_policy: Union[str, SegregationPolicy] = SegregationPolicy.CIOC,
    ciocs_per_cabinet: int = 4,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> Union[PipelineResult, Tuple[List[AllocatedSignal], bytes, List[TechnicalQuery]]]:
    ...


def run_hierarchical_pipeline(
    file_or_path_or_bytes: Union[str, Path, BinaryIO, bytes],
    charms_per_baseplate: int = 12,
    baseplates_per_cioc: int = 8,
    spare_margin: float = 0.0,
    declared_sizes: Optional[dict[str, int]] = None,
    catalog: Optional[Dict[str, Any]] = None,
    return_tqs: bool = False,
    segregation_policy: Union[str, SegregationPolicy] = SegregationPolicy.CIOC,
    ciocs_per_cabinet: int = 4,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
) -> Union[PipelineResult, Tuple[List[AllocatedSignal], bytes, List[TechnicalQuery]]]:
    """Execute end-to-end two-stage hierarchical packing, independent audit, and Excel export.

    Args:
        file_or_path_or_bytes: Path to input spreadsheet, binary buffer, or raw bytes.
        charms_per_baseplate: Capacity of each physical baseplate (12).
        baseplates_per_cioc: Number of baseplates per CIOC card (8).
        spare_margin: Reserve capacity margin per CIOC (e.g. 0.0).
        declared_sizes: Optional dictionary of cable_id -> pair capacity.
        catalog: Optional custom hardware catalog mapping.
        return_tqs: If True, returns explicit 3-tuple (allocations, excel_bytes, tqs).
        segregation_policy: "CIOC" (CIOC-level segregation) or "CABINET" (Cabinet-level segregation).
        ciocs_per_cabinet: Number of CIOCs per Marshalling Cabinet (default: 4).
        template_file_or_bytes: Optional blank customer Excel template to overlay onto.

    Returns:
        PipelineResult (unpackable as (allocations, excel_bytes) or providing .tqs).
    """
    if isinstance(file_or_path_or_bytes, bytes):
        buffer: Union[str, Path, BinaryIO] = io.BytesIO(file_or_path_or_bytes)
    else:
        buffer = file_or_path_or_bytes

    # 1. Parse and validate with fault isolation
    signals, tqs = parse_and_validate_signals(buffer)

    # 2. Solve Two-Stage Hierarchical Packing
    if signals:
        allocations = solve_hierarchical_packing(
            signals,
            declared_sizes=declared_sizes,
            charms_per_baseplate=charms_per_baseplate,
            baseplates_per_cioc=baseplates_per_cioc,
            spare_margin=spare_margin,
            catalog=catalog,
            segregation_policy=segregation_policy,
            ciocs_per_cabinet=ciocs_per_cabinet,
        )

        # 3. Independent Post-Solve Hard Constraint Audit
        valid, errors = verify_allocation_invariants(
            allocations,
            charms_per_baseplate=charms_per_baseplate,
            baseplates_per_cioc=baseplates_per_cioc,
            segregation_policy=segregation_policy,
            ciocs_per_cabinet=ciocs_per_cabinet,
        )
        if not valid:
            raise RuntimeError(f"Allocation violated physical constraints: {'; '.join(errors)}")
    else:
        allocations = []

    # 4. Export populated Excel workbook with dedicated TQ_Log and all deliverable sheets
    excel_bytes = export_allocation_to_excel(
        allocations,
        tqs=tqs,
        template_path_or_buffer=template_file_or_bytes,
        charms_per_bp=charms_per_baseplate,
        bps_per_cioc=baseplates_per_cioc,
        ciocs_per_cabinet=ciocs_per_cabinet,
        segregation_policy=segregation_policy,
    )

    if return_tqs:
        return allocations, excel_bytes, tqs

    return PipelineResult(allocations, excel_bytes, tqs)


def _parse_allocated_signals_from_excel(file_or_path_or_bytes: Any) -> Optional[List[AllocatedSignal]]:
    """Try to read allocated signals directly if buffer is an exported deliverable."""
    import openpyxl

    try:
        if isinstance(file_or_path_or_bytes, bytes):
            buf = io.BytesIO(file_or_path_or_bytes)
        elif hasattr(file_or_path_or_bytes, "read"):
            buf = file_or_path_or_bytes
        elif isinstance(file_or_path_or_bytes, (str, Path)):
            with open(file_or_path_or_bytes, "rb") as f:
                buf = io.BytesIO(f.read())
        else:
            return None

        wb = openpyxl.load_workbook(buf, data_only=True)
        sheet_name = "Allocation" if "Allocation" in wb.sheetnames else None
        if not sheet_name:
            return None
        ws = wb[sheet_name]
        headers = [str(cell.value or "").strip().lower() for cell in ws[1]]
        if "slot" not in headers or "baseplate" not in headers:
            return None
        col_map = {name: idx for idx, name in enumerate(headers, start=1)}
        allocs = []
        for row in range(2, ws.max_row + 1):
            tag_val = str(ws.cell(row=row, column=col_map.get("tag", 1)).value or "").strip()
            if not tag_val:
                continue
            io_str = str(ws.cell(row=row, column=col_map.get("io type", 2)).value or "AI").strip().upper()
            bp_val = int(ws.cell(row=row, column=col_map.get("baseplate", 5)).value or 1)
            slot_val = int(ws.cell(row=row, column=col_map.get("slot", 6)).value or 1)
            cioc_val = int(ws.cell(row=row, column=col_map.get("cioc", 4)).value or 1)
            cab_cell = ws.cell(row=row, column=col_map.get("cabinet", 3)).value
            if isinstance(cab_cell, str) and "CAB-" in cab_cell.upper():
                cab_val = int(cab_cell.upper().replace("CAB-", ""))
            else:
                cab_val = int(cab_cell or 1)

            cbl_val = ws.cell(row=row, column=col_map.get("cable id", 7)).value
            p_val = ws.cell(row=row, column=col_map.get("pair no", 8)).value
            seg_val = ws.cell(row=row, column=col_map.get("segregation", 13)).value
            desc_val = ws.cell(row=row, column=col_map.get("description", 14)).value
            status_val = str(ws.cell(row=row, column=col_map.get("status", 15)).value or "")

            allocs.append(
                AllocatedSignal(
                    tag=tag_val,
                    io_type=IOType(io_str) if io_str in IOType.__members__ else IOType.AI,
                    cabinet=cab_val,
                    cioc=cioc_val,
                    baseplate=bp_val,
                    slot=slot_val,
                    cable_id=str(cbl_val) if cbl_val and str(cbl_val) != "-" else None,
                    pair_no=int(p_val) if p_val and str(p_val) != "-" else None,
                    segregation=str(seg_val).strip().upper() if seg_val and str(seg_val) in ("A", "B") else None,
                    description=str(desc_val) if desc_val else None,
                    is_spare="SPARE" in status_val.upper(),
                )
            )
        return allocs if allocs else None
    except Exception:
        return None


def run_revision_pipeline(
    rev_a_file_or_data: Union[str, Path, BinaryIO, bytes, Sequence[AllocatedSignal]],
    rev_b_file_or_data: Union[str, Path, BinaryIO, bytes, Sequence[Signal]],
    charms_per_baseplate: int = 12,
    baseplates_per_cioc: int = 8,
    ciocs_per_cabinet: int = 4,
    segregation_policy: Union[str, SegregationPolicy] = SegregationPolicy.CIOC,
    catalog: Optional[Dict[str, Any]] = None,
    template_file_or_bytes: Optional[Union[str, Path, BinaryIO, bytes]] = None,
    time_limit_seconds: float = 10.0,
) -> RevisionPipelineResult:
    """Execute end-to-end customer revision diff, perturbation minimization, audit, and deliverable export.

    Args:
        rev_a_file_or_data: Baseline Rev A spreadsheet, workbook, or sequence of AllocatedSignals.
        rev_b_file_or_data: New Rev B spreadsheet, buffer, or sequence of Signals.
        charms_per_baseplate: Slots per baseplate (12).
        baseplates_per_cioc: Baseplates per CIOC (8).
        ciocs_per_cabinet: CIOCs per Marshalling Cabinet (4).
        segregation_policy: "CIOC" or "CABINET" segregation rule.
        catalog: Optional custom DeltaV hardware catalog.
        template_file_or_bytes: Optional customer Excel template to overlay onto.
        time_limit_seconds: Solver maximum run time.

    Returns:
        RevisionPipelineResult containing final allocations, Excel deliverable with Revision_Impact tab, and diff metrics.
    """
    from src.revision.solver import solve_revision_allocation

    tqs_a: List[TechnicalQuery] = []
    tqs_b: List[TechnicalQuery] = []

    # 1. Obtain Rev A baseline allocations
    if (
        isinstance(rev_a_file_or_data, (list, tuple))
        and rev_a_file_or_data
        and isinstance(rev_a_file_or_data[0], AllocatedSignal)
    ):
        rev_a_allocations = list(rev_a_file_or_data)
    else:
        # Check if rev_a already contains allocated deliverable sheet
        existing_allocs = _parse_allocated_signals_from_excel(rev_a_file_or_data)
        if existing_allocs:
            rev_a_allocations = existing_allocs
        else:
            # Parse and solve baseline Rev A
            if isinstance(rev_a_file_or_data, bytes):
                buf_a: Union[str, Path, BinaryIO] = io.BytesIO(rev_a_file_or_data)
            elif isinstance(rev_a_file_or_data, (str, Path, io.IOBase)):
                buf_a = rev_a_file_or_data
            else:
                raise TypeError(f"Unsupported rev_a input type: {type(rev_a_file_or_data)}")
            sigs_a, tqs_a = parse_and_validate_signals(buf_a)
            rev_a_allocations = solve_hierarchical_packing(
                sigs_a,
                charms_per_baseplate=charms_per_baseplate,
                baseplates_per_cioc=baseplates_per_cioc,
                ciocs_per_cabinet=ciocs_per_cabinet,
                segregation_policy=segregation_policy,
                catalog=catalog,
                time_limit_seconds=time_limit_seconds,
            )

    # 2. Obtain Rev B signals
    if (
        isinstance(rev_b_file_or_data, (list, tuple))
        and rev_b_file_or_data
        and isinstance(rev_b_file_or_data[0], Signal)
    ):
        rev_b_signals = list(rev_b_file_or_data)
    else:
        if isinstance(rev_b_file_or_data, bytes):
            buf_b: Union[str, Path, BinaryIO] = io.BytesIO(rev_b_file_or_data)
        elif isinstance(rev_b_file_or_data, (str, Path, io.IOBase)):
            buf_b = rev_b_file_or_data
        else:
            raise TypeError(f"Unsupported rev_b input type: {type(rev_b_file_or_data)}")
        rev_b_signals, tqs_b = parse_and_validate_signals(buf_b)

    # 3. Solve revision perturbation minimization & freezing
    rev_result = solve_revision_allocation(
        rev_a_allocations=rev_a_allocations,
        rev_b_signals=rev_b_signals,
        charms_per_baseplate=charms_per_baseplate,
        baseplates_per_cioc=baseplates_per_cioc,
        ciocs_per_cabinet=ciocs_per_cabinet,
        segregation_policy=segregation_policy,
        catalog=catalog,
        time_limit_seconds=time_limit_seconds,
    )

    # 4. Independent Audit Pass
    valid, errors = verify_allocation_invariants(
        rev_result.allocations,
        charms_per_baseplate=charms_per_baseplate,
        baseplates_per_cioc=baseplates_per_cioc,
        ciocs_per_cabinet=ciocs_per_cabinet,
        segregation_policy=segregation_policy,
    )
    if not valid:
        raise RuntimeError(f"Revision allocation violated physical constraints: {'; '.join(errors)}")

    # 5. Export Deliverable Workbook with Revision_Impact tab
    all_tqs = tqs_a + tqs_b
    excel_bytes = export_allocation_to_excel(
        rev_result.allocations,
        tqs=all_tqs,
        template_path_or_buffer=template_file_or_bytes,
        charms_per_bp=charms_per_baseplate,
        bps_per_cioc=baseplates_per_cioc,
        ciocs_per_cabinet=ciocs_per_cabinet,
        segregation_policy=segregation_policy,
        revision_result=rev_result,
        rev_a_allocations=rev_a_allocations,
    )

    return RevisionPipelineResult(
        allocations=rev_result.allocations,
        excel_bytes=excel_bytes,
        revision_result=rev_result,
        tqs_a=tqs_a,
        tqs_b=tqs_b,
    )
