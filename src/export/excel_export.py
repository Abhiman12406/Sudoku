"""Excel export module for allocated DeltaV signals."""

import io
from pathlib import Path
from typing import BinaryIO, Optional, Sequence, Union
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.domain.models import AllocatedSignal

COLUMNS = [
    ("Tag", "tag"),
    ("IO Type", "io_type"),
    ("Baseplate", "baseplate"),
    ("Slot", "slot"),
    ("Cable ID", "cable_id"),
    ("Pair No", "pair_no"),
    ("CHARM Part No", "charm_part_no"),
    ("Terminal Block Type", "terminal_block_type"),
    ("Description", "description"),
]


def export_allocation_to_excel(
    allocations: Sequence[AllocatedSignal],
    output_path_or_buffer: Optional[Union[str, Path, BinaryIO]] = None,
) -> bytes:
    """Export allocated signals to a styled Excel workbook (.xlsx).

    Args:
        allocations: Sequence of AllocatedSignal objects.
        output_path_or_buffer: Optional file path or buffer to write to.

    Returns:
        Raw bytes of the generated Excel workbook.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Allocation"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Write headers
    for col_idx, (col_name, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Write rows
    center_alignment = Alignment(horizontal="center", vertical="center")
    left_alignment = Alignment(horizontal="left", vertical="center")

    for row_idx, alloc in enumerate(allocations, start=2):
        for col_idx, (_, attr_name) in enumerate(COLUMNS, start=1):
            val = getattr(alloc, attr_name)
            # Normalize IOType enum to string
            if hasattr(val, "value"):
                val = val.value
            cell = ws.cell(row=row_idx, column=col_idx, value=val)

            if attr_name in ("baseplate", "slot", "pair_no", "io_type"):
                cell.alignment = center_alignment
            else:
                cell.alignment = left_alignment

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Freeze panes on header row
    ws.freeze_panes = "A2"

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()

    if output_path_or_buffer is not None:
        if isinstance(output_path_or_buffer, (str, Path)):
            with open(output_path_or_buffer, "wb") as f:
                f.write(excel_bytes)
        elif hasattr(output_path_or_buffer, "write"):
            output_path_or_buffer.write(excel_bytes)

    return excel_bytes
