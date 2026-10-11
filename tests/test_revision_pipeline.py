"""Unit and E2E tests for Seam 4: Revision Pipeline."""

import io
import openpyxl
import pandas as pd
import pytest

from src.pipeline import RevisionPipelineResult, run_revision_pipeline


def _create_sample_excel_bytes(data: dict) -> bytes:
    df = pd.DataFrame(data)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Signal_Schedule")
    return buf.getvalue()


def test_run_revision_pipeline_e2e():
    # Rev A baseline
    rev_a_data = {
        "Tag": ["10-FT-101", "10-PT-102", "10-TT-103", "10-XV-104"],
        "IO Type": ["AI", "AI", "RTD", "DO"],
        "Cable ID": ["CBL-01", "CBL-01", "CBL-02", "CBL-02"],
        "Pair No": [1, 2, 1, 2],
        "Segregation": ["A", "A", "A", "A"],
        "Description": ["Flow", "Pressure", "Temp", "Valve"],
    }
    rev_a_bytes = _create_sample_excel_bytes(rev_a_data)

    # Rev B:
    # 10-FT-101: unchanged
    # 10-PT-102: modified (IO Type changed from AI to AO)
    # 10-TT-103: deleted (not in Rev B)
    # 10-XV-104: unchanged
    # 10-LT-105: added (new in Rev B)
    rev_b_data = {
        "Tag": ["10-FT-101", "10-PT-102", "10-XV-104", "10-LT-105"],
        "IO Type": ["AI", "AO", "DO", "AI"],
        "Cable ID": ["CBL-01", "CBL-01", "CBL-02", "CBL-03"],
        "Pair No": [1, 2, 2, 1],
        "Segregation": ["A", "A", "A", "A"],
        "Description": ["Flow", "Pressure Valve", "Valve", "Level"],
    }
    rev_b_bytes = _create_sample_excel_bytes(rev_b_data)

    result = run_revision_pipeline(rev_a_bytes, rev_b_bytes)

    assert isinstance(result, RevisionPipelineResult)
    assert len(result.allocations) >= 4  # Rev B signals allocated
    assert result.revision_result is not None

    diff = result.revision_result.diff_result
    assert len(diff.unchanged) == 2  # 10-FT-101, 10-XV-104
    assert len(diff.modified) == 1   # 10-PT-102
    assert len(diff.deleted) == 1    # 10-TT-103
    assert len(diff.added) == 1      # 10-LT-105

    # Check Excel workbook contains Revision_Impact
    wb = openpyxl.load_workbook(io.BytesIO(result.excel_bytes))
    assert "Revision_Impact" in wb.sheetnames
    assert "Allocation" in wb.sheetnames
    assert "Rule_Compliance" in wb.sheetnames
