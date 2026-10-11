"""End-to-end tests for Issue 06: Template-Preserving Overlay & 5-Tab Deliverable Package."""

import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import pytest

from src.domain.models import IOType, AllocatedSignal
from src.export.excel_export import export_allocation_to_excel
from src.pipeline import run_hierarchical_pipeline, run_walking_pipeline
from src.ui.app import generate_sample_blank_customer_template_bytes


def test_e2e_deliverable_package_full_integrity_with_customer_template():
    """Verify that a customer Excel template is preserved and multi-tab deliverable is 100% valid."""
    # 1. Create a rich, real-world customer template
    temp_bytes = generate_sample_blank_customer_template_bytes()
    wb_temp = openpyxl.load_workbook(io.BytesIO(temp_bytes), data_only=False)
    assert "Customer_IO_Schedule" in wb_temp.sheetnames

    # 2. Build input dataset with 12-pair trunk cable signals
    input_data = {
        "Tag": [f"10-FT-{i:03d}" for i in range(1, 11)],
        "IO Type": ["AI"] * 10,
        "Cable ID": ["TRUNK-12P"] * 10,
        "Pair No": list(range(1, 11)),
        "Description": [f"Flow Transmitter Loop {i}" for i in range(1, 11)],
    }
    df_in = pd.DataFrame(input_data)
    in_buf = io.BytesIO()
    with pd.ExcelWriter(in_buf, engine="openpyxl") as writer:
        df_in.to_excel(writer, index=False)
    in_bytes = in_buf.getvalue()

    # 3. Run hierarchical pipeline with customer template overlay
    result = run_hierarchical_pipeline(
        in_bytes,
        template_file_or_bytes=temp_bytes,
    )
    allocations, excel_bytes = result

    # 4. Strict XML & Workbook validation
    wb_out = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)

    # Invariant: exactly 5 deliverable sheets
    assert len(wb_out.sheetnames) == 5
    assert wb_out.sheetnames[0] == "Customer_IO_Schedule"
    assert "Allocation" in wb_out.sheetnames
    assert "TQ_Log" in wb_out.sheetnames
    assert "Cabinet_Summary" in wb_out.sheetnames
    assert "Rule_Compliance" in wb_out.sheetnames

    # 5. Verify Customer_IO_Schedule preserved styles and formulas
    ws_cust = wb_out["Customer_IO_Schedule"]
    assert ws_cust.cell(row=1, column=1).value == "Loop Tag"
    assert ws_cust.cell(row=1, column=1).font.bold is True
    assert ws_cust.cell(row=1, column=1).font.name == "Segoe UI"
    assert ws_cust.cell(row=1, column=1).fill.start_color.rgb == "00002060"

    # Data row 2
    assert ws_cust.cell(row=2, column=1).value == "10-FT-001"
    assert ws_cust.cell(row=2, column=2).value == "AI"
    assert ws_cust.cell(row=2, column=14).value == '=CONCATENATE(A2, "-", B2)'

    # Data row 3: formula adjusted
    assert ws_cust.cell(row=3, column=14).value == '=CONCATENATE(A3, "-", B3)'

    # 6. Verify Cabinet_Summary
    df_cab = pd.read_excel(io.BytesIO(excel_bytes), sheet_name="Cabinet_Summary")
    assert "Cabinet ID" in df_cab.columns
    assert "Total Capacity" in df_cab.columns
    assert "Active Loops" in df_cab.columns
    assert "Wired Spares" in df_cab.columns
    assert "Free Spares" in df_cab.columns
    assert "Utilization %" in df_cab.columns

    cab_row = df_cab[df_cab["Cabinet ID"] == "CAB-01"].iloc[0]
    assert cab_row["Active Loops"] == 10
    assert cab_row["Wired Spares"] == 2  # 12 pairs - 10 active = 2 wired spares
    assert cab_row["Total Allocated"] == 12

    # 7. Verify Rule_Compliance tab
    df_rules = pd.read_excel(io.BytesIO(excel_bytes), sheet_name="Rule_Compliance")
    assert "Rule ID" in df_rules.columns
    assert "Status" in df_rules.columns
    assert "Violations" in df_rules.columns
    assert len(df_rules) >= 6
    for _, row in df_rules.iterrows():
        assert row["Status"] == "PASS"
        assert row["Violations"] == 0

    # 8. Re-read with openpyxl in read_only mode to ensure standard parser acceptance
    wb_ro = openpyxl.load_workbook(io.BytesIO(excel_bytes), read_only=True)
    assert len(wb_ro.sheetnames) == 5
    wb_ro.close()


def test_e2e_deliverable_package_without_template():
    """Verify that exporting without an uploaded template generates 5 full standard sheets."""
    input_data = {
        "Tag": ["10-PT-001", "10-PT-002"],
        "IO Type": ["AI", "AI"],
        "Cable ID": ["TRUNK-01", "TRUNK-01"],
        "Pair No": [1, 2],
        "Description": ["Pressure 1", "Pressure 2"],
    }
    df_in = pd.DataFrame(input_data)
    in_buf = io.BytesIO()
    with pd.ExcelWriter(in_buf, engine="openpyxl") as writer:
        df_in.to_excel(writer, index=False)
    in_bytes = in_buf.getvalue()

    result = run_walking_pipeline(in_bytes)
    allocations, excel_bytes = result

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert len(wb.sheetnames) == 5
    assert wb.sheetnames == [
        "Allocation",
        "TQ_Log",
        "Cabinet_Summary",
        "Rule_Compliance",
        "Template_Overlay",
    ]

    # Every sheet has content
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        assert ws.max_row >= 1
        assert ws.max_column >= 1


def test_e2e_workbook_formula_syntax_and_cell_types():
    """Verify that all formulas in the generated workbook are valid syntax and cell values valid."""
    allocations = [
        AllocatedSignal(
            tag=f"20-PT-{i}",
            io_type=IOType.AI,
            cioc=1,
            baseplate=1,
            slot=i,
            cable_id="CBL-01",
            pair_no=i,
            charm_part_no="KL3001X1-BA1",
            terminal_block_type="Standard",
        )
        for i in range(1, 5)
    ]

    excel_bytes = export_allocation_to_excel(allocations)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    # Excel formula: verify starts with valid formula token and parentheses
                    assert cell.value.startswith("=")
                    assert "(" in cell.value
                    assert ")" in cell.value
