import io
import pandas as pd
import pytest

from src.ui.app import (
    catalog_to_dataframe,
    dataframe_to_catalog,
    format_severity_badge,
    generate_corrupt_sample_template_bytes,
    generate_sample_template_bytes,
)


def test_format_severity_badge():
    assert format_severity_badge("Critical") == "🔴 Critical"
    assert format_severity_badge("Warning") == "🟡 Warning"
    assert format_severity_badge("Info") == "🔵 Info"
    assert format_severity_badge("CRITICAL") == "🔴 Critical"


def test_generate_sample_template_bytes():
    clean_bytes = generate_sample_template_bytes()
    assert isinstance(clean_bytes, bytes)
    df = pd.read_excel(io.BytesIO(clean_bytes))
    assert len(df) == 14
    assert "Tag" in df.columns
    assert "IO Type" in df.columns


def test_generate_corrupt_sample_template_bytes():
    corrupt_bytes = generate_corrupt_sample_template_bytes()
    assert isinstance(corrupt_bytes, bytes)
    df = pd.read_excel(io.BytesIO(corrupt_bytes))
    assert len(df) >= 7
    assert "XYZ_CORRUPT" in list(df["IO Type"])


def test_generate_sample_blank_customer_template_bytes():
    import openpyxl
    from src.ui.app import generate_sample_blank_customer_template_bytes

    temp_bytes = generate_sample_blank_customer_template_bytes()
    assert isinstance(temp_bytes, bytes)
    wb = openpyxl.load_workbook(io.BytesIO(temp_bytes), data_only=False)
    ws = wb.active
    assert ws.title == "Customer_IO_Schedule"
    assert ws.cell(row=1, column=1).value == "Loop Tag"
    # Verify template formula in row 2
    assert str(ws.cell(row=2, column=14).value).startswith("=CONCATENATE")



def test_catalog_to_dataframe_and_roundtrip():
    test_catalog = {
        "AI": {
            "charm_description": "Analog Input 4-20 mA HART",
            "charm_part_no": "KL3001X1-BA1",
            "terminal_block_type": "Standard",
            "terminal_block_part_no": "KL4001X1-BA1",
        },
        "TC": {
            "charm_description": "Thermocouple / mV Input",
            "charm_part_no": "KL3006X1-BA1",
            "terminal_block_type": "CJC",
            "terminal_block_part_no": "KL4004X1-BA1",
        },
    }

    df = catalog_to_dataframe(test_catalog)
    assert isinstance(df, pd.DataFrame)
    assert len(df) == 8
    assert set(df["Signal Type"]) == {"AI", "AO", "DI", "DO", "DO_RELAY", "TC", "RTD", "SPARE"}

    # Modify and convert back
    df.loc[df["Signal Type"] == "TC", "CHARM Part No"] = "KL3006X1-MOD"
    reconstructed = dataframe_to_catalog(df)
    assert reconstructed["TC"]["charm_part_no"] == "KL3006X1-MOD"
    assert reconstructed["TC"]["terminal_block_type"] == "CJC"


def test_generate_sample_rev_a_and_b_bytes():
    from src.ui.app import generate_sample_rev_a_bytes, generate_sample_rev_b_bytes

    a_bytes = generate_sample_rev_a_bytes()
    b_bytes = generate_sample_rev_b_bytes()

    assert isinstance(a_bytes, bytes)
    assert isinstance(b_bytes, bytes)

    df_a = pd.read_excel(io.BytesIO(a_bytes))
    df_b = pd.read_excel(io.BytesIO(b_bytes))

    assert len(df_a) > 0
    assert len(df_b) > 0
    assert "Tag" in df_a.columns
    assert "Tag" in df_b.columns


def test_highlight_diff_dataframe():
    from src.ui.app import highlight_diff_dataframe

    test_data = [
        {"Status": "🟢 ADDED", "Tag": "10-LT-101", "Displacement": "NEW ALLOCATION", "Penalty Δ": 0},
        {"Status": "🔴 DELETED", "Tag": "10-PT-102", "Displacement": "DELETED (Vacated)", "Penalty Δ": 0},
        {"Status": "🟠 MODIFIED", "Tag": "10-FT-103", "Displacement": "RELOCATED", "Penalty Δ": 1000},
        {"Status": "⚪ UNCHANGED", "Tag": "10-TT-104", "Displacement": "STATIONARY (Frozen)", "Penalty Δ": 0},
    ]
    df = pd.DataFrame(test_data)
    styler = highlight_diff_dataframe(df)
    assert hasattr(styler, "to_html") or hasattr(styler, "style")


