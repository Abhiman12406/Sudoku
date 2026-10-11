"""Unit tests for Seam 3: Revision Impact Deliverable Export."""

import io
import openpyxl
import pytest

from src.domain.models import AllocatedSignal, IOType, Signal
from src.export.excel_export import export_allocation_to_excel
from src.revision.classifier import classify_revision_diff
from src.revision.excel_tab import populate_revision_impact_sheet
from src.revision.solver import RevisionAllocationResult, solve_revision_allocation


def test_revision_impact_sheet_generation():
    rev_a = [
        AllocatedSignal(
            tag="10-FT-101",
            io_type=IOType.AI,
            cabinet=1,
            cioc=1,
            baseplate=1,
            slot=1,
        ),
        AllocatedSignal(
            tag="10-PT-102",
            io_type=IOType.AI,
            cabinet=1,
            cioc=1,
            baseplate=1,
            slot=2,
        ),
    ]

    rev_b = [
        Signal(tag="10-FT-101", io_type=IOType.AI),  # UNCHANGED
        Signal(tag="10-LT-103", io_type=IOType.AI),  # ADDED (backfills vacated slot 2)
    ]
    # PT-102 is DELETED

    rev_result = solve_revision_allocation(rev_a, rev_b)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revision_Impact"
    populate_revision_impact_sheet(ws, rev_result, rev_a)

    # Verify sheet contents
    all_values = []
    for row in ws.iter_rows(values_only=True):
        all_values.extend([str(v) for v in row if v is not None])

    content_str = " ".join(all_values)

    # Check KPI metrics
    assert "Revision Impact Summary" in content_str
    assert "Unchanged Loops" in content_str
    assert "Added Loops" in content_str
    assert "Deleted Loops" in content_str

    # Check signal tags in delta table
    assert "10-FT-101" in content_str
    assert "10-PT-102" in content_str
    assert "10-LT-103" in content_str

    # Check status strings
    assert "UNCHANGED" in content_str
    assert "DELETED" in content_str
    assert "ADDED" in content_str


def test_export_allocation_to_excel_with_revision_result():
    rev_a = [
        AllocatedSignal(
            tag="10-FT-101",
            io_type=IOType.AI,
            cabinet=1,
            cioc=1,
            baseplate=1,
            slot=1,
        )
    ]
    rev_b = [
        Signal(tag="10-FT-101", io_type=IOType.AI),
        Signal(tag="10-LT-102", io_type=IOType.AO),
    ]
    rev_result = solve_revision_allocation(rev_a, rev_b)

    excel_bytes = export_allocation_to_excel(
        rev_result.allocations,
        revision_result=rev_result,
        rev_a_allocations=rev_a,
    )

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Revision_Impact" in wb.sheetnames
    assert "Allocation" in wb.sheetnames
    assert "Rule_Compliance" in wb.sheetnames
